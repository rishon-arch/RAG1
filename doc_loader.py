import os
import re
import hashlib
from typing import List, Callable, Optional, Dict, Any
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader, CSVLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".txt", ".md", ".csv")

# Friendly icon mapping for Streamlit UI
FILE_TYPE_ICONS = {
    ".pdf": "📄",
    ".docx": "📝",
    ".pptx": "📊",
    ".txt": "📑",
    ".md": "📋",
    ".csv": "📈"
}

def get_collection_name(file_path: str) -> str:
    """
    Generate a deterministic, valid Chroma collection name for a given file.
    Preserves 'langchain' for 'nptel.pdf' to maintain the existing pre-indexed database.
    Chroma requirements:
      - 3-63 characters
      - Alphanumeric, underscores, hyphens
      - Starts and ends with an alphanumeric character
    """
    base = os.path.basename(file_path)
    if base.lower() == "nptel.pdf":
        return "langchain"

    clean = re.sub(r'[^a-zA-Z0-9_-]', '_', os.path.splitext(base)[0]).strip('_')
    if not clean:
        clean = "doc"
    clean = clean[:35]

    # Deterministic 8-char hash based on normalized filename
    h = hashlib.md5(base.lower().encode("utf-8")).hexdigest()[:8]
    col_name = f"doc_{clean}_{h}"
    return col_name[:63]

def format_citation(metadata: dict) -> str:
    """Format human-friendly citation from chunk metadata (page, slide, row, or section)."""
    page = metadata.get("page")
    if page is not None:
        if isinstance(page, int):
            return f"Page {page + 1}"
        return f"{page}"
    if "row" in metadata:
        return f"Row {metadata['row']}"
    return "Document"

def load_raw_documents(file_path: str, progress_callback: Optional[Callable[[int, str], None]] = None) -> List[Document]:
    """
    Load raw documents based on file extension with appropriate metadata.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: '{file_path}'")

    ext = os.path.splitext(file_path)[1].lower()
    docs = []

    if ext == ".pdf":
        loader = PyPDFLoader(file_path)
        for idx, page in enumerate(loader.lazy_load()):
            if "page" not in page.metadata:
                page.metadata["page"] = idx
            docs.append(page)
            if progress_callback and (idx + 1) % 100 == 0:
                progress_callback(idx + 1, f"Extracted {idx + 1} PDF pages...")

    elif ext == ".docx":
        loader = Docx2txtLoader(file_path)
        raw_docs = loader.load()
        for doc in raw_docs:
            doc.metadata["page"] = "Document"
            docs.append(doc)

    elif ext == ".pptx":
        prs = Presentation(file_path)
        for idx, slide in enumerate(prs.slides, start=1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        t = paragraph.text.strip()
                        if t:
                            slide_texts.append(t)
            if slide_texts:
                docs.append(Document(
                    page_content="\n".join(slide_texts),
                    metadata={"source": file_path, "page": f"Slide {idx}"}
                ))

    elif ext in (".txt", ".md"):
        encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
        content = None
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except Exception:
                continue
        if content is None:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        docs.append(Document(
            page_content=content,
            metadata={"source": file_path, "page": "Section"}
        ))

    elif ext == ".csv":
        try:
            loader = CSVLoader(file_path, encoding="utf-8")
            raw_docs = loader.load()
        except Exception:
            loader = CSVLoader(file_path, encoding="latin-1")
            raw_docs = loader.load()
        for d in raw_docs:
            row_num = d.metadata.get("row", d.metadata.get("line", "Unknown"))
            d.metadata["page"] = f"Row {row_num}"
            docs.append(d)

    else:
        raise ValueError(f"Unsupported file format: '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}")

    return docs

def split_documents(docs: List[Document], chunk_size: int = 1200, chunk_overlap: int = 200) -> List[Document]:
    """
    Split loaded documents into overlapping chunks suitable for embedding.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap
    )
    return splitter.split_documents(docs)

def get_available_documents(search_dirs: Optional[List[str]] = None) -> List[dict]:
    """
    Find all documents with supported extensions in the project.
    Returns a list of dicts: {'path': ..., 'name': ..., 'size_mb': ..., 'ext': ..., 'icon': ...}
    """
    if search_dirs is None:
        search_dirs = [".", "./uploaded_docs"]

    found = {}
    ignored_dirs = {".venv", "venv", "__pycache__", "chroma_local_db", ".git", ".gemini", "scratch"}
    ignored_files = {"requirements.txt"}

    for d in search_dirs:
        if not os.path.exists(d):
            continue
        for root, dirs, files in os.walk(d):
            dirs[:] = [sub for sub in dirs if sub not in ignored_dirs and not sub.startswith(".")]
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in SUPPORTED_EXTENSIONS and file not in ignored_files:
                    full_path = os.path.normpath(os.path.join(root, file))
                    rel_path = os.path.relpath(full_path, ".").replace("\\", "/")
                    if rel_path not in found:
                        size_mb = os.path.getsize(full_path) / (1024 * 1024)
                        found[rel_path] = {
                            "path": rel_path,
                            "name": file,
                            "size_mb": round(size_mb, 2),
                            "ext": ext,
                            "icon": FILE_TYPE_ICONS.get(ext, "📄")
                        }

    # Sort: nptel.pdf first (default benchmark), then alphabetical
    def sort_key(item):
        if item["name"].lower() == "nptel.pdf":
            return (0, item["name"])
        return (1, item["name"].lower())

    return sorted(list(found.values()), key=sort_key)

def federated_query(
    client,
    collection_names: List[str],
    query_text: str,
    embeddings,
    k_per_col: int = 3,
    total_k: int = 5
) -> List[Dict[str, Any]]:
    """
    Performs a single-pass embedding on the query, queries each specified
    Chroma collection, aggregates hits, sorts by distance (ascending),
    and returns the top total_k hits.
    
    Each hit contains:
      - 'content': str
      - 'metadata': dict
      - 'distance': float
      - 'collection': str
    """
    if not collection_names:
        return []

    # Single pass embedding on GPU via FastEmbed
    query_vector = embeddings.embed_query(query_text)
    all_hits = []

    for col_name in collection_names:
        try:
            col = client.get_collection(col_name)
            count = col.count()
            if count == 0:
                continue
            actual_k = min(k_per_col, count)
            res = col.query(query_embeddings=[query_vector], n_results=actual_k)
            if res and res.get("documents") and res["documents"][0]:
                docs = res["documents"][0]
                metas = res["metadatas"][0] if res.get("metadatas") else [{}] * len(docs)
                dists = res["distances"][0] if res.get("distances") else [0.0] * len(docs)
                for d, m, dist in zip(docs, metas, dists):
                    all_hits.append({
                        "content": d,
                        "metadata": m or {},
                        "distance": dist,
                        "collection": col_name
                    })
        except Exception:
            continue

    # Sort ascending by distance (lowest distance = highest similarity)
    all_hits.sort(key=lambda x: x["distance"])
    return all_hits[:total_k]
